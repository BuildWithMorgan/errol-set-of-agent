import asyncio
import io
import json
import os
import uuid
from contextlib import asynccontextmanager
from datetime import datetime
from pathlib import Path
from typing import AsyncGenerator, Optional

import httpx
from docx import Document
from docx.shared import Pt, RGBColor
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from fastapi import FastAPI, HTTPException, Request, UploadFile, File
from fastapi.responses import StreamingResponse, JSONResponse, Response
from fastapi.staticfiles import StaticFiles
from pypdf import PdfReader

from dotenv import load_dotenv
load_dotenv()

import tools.file_cleaner as file_cleaner
from prompts import build_prompt, build_input_summary, AGENT_LABELS

_history_lock = asyncio.Lock()

# ─── Config ───────────────────────────────────────────────────────────────────

OLLAMA_URL = os.getenv("OLLAMA_BASE_URL", "http://localhost:11434")
OLLAMA_MODEL = os.getenv("OLLAMA_MODEL", "gemma4:e4b")

DATA_DIR = Path("data")
HISTORY_FILE = DATA_DIR / "history.json"
TEMPLATES_FILE = DATA_DIR / "templates.json"
FEEDBACK_FILE = DATA_DIR / "feedback.json"
CLEANER_SETTINGS_FILE = DATA_DIR / "cleaner_settings.json"


@asynccontextmanager
async def lifespan(app_: FastAPI):
    yield


app = FastAPI(lifespan=lifespan)

# ─── JSON helpers ─────────────────────────────────────────────────────────────

def read_json(path: Path) -> list:
    if not path.exists():
        return []
    return json.loads(path.read_text(encoding="utf-8"))


def write_json(path: Path, data: list) -> None:
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")


# ─── Cleaner helpers ──────────────────────────────────────────────────────────

def get_cleaner_settings() -> dict:
    defaults = {
        "folders": os.getenv(
            "CLEANER_FOLDERS",
            f"{Path.home()}/Desktop,{Path.home()}/Downloads"
        ).split(","),
    }
    if CLEANER_SETTINGS_FILE.exists():
        saved = json.loads(CLEANER_SETTINGS_FILE.read_text(encoding="utf-8"))
        defaults.update(saved)
    return defaults


# ─── Status ───────────────────────────────────────────────────────────────────

@app.get("/api/status")
async def status():
    try:
        async with httpx.AsyncClient(timeout=3) as client:
            r = await client.get(f"{OLLAMA_URL}/api/tags")
            if r.status_code == 200:
                return {"online": True, "model": OLLAMA_MODEL}
    except Exception:
        pass
    return {"online": False, "model": OLLAMA_MODEL}

# ─── File upload & text extraction ───────────────────────────────────────────

@app.post("/api/upload")
async def upload_file(file: UploadFile = File(...)):
    filename = file.filename or ""
    content_type = file.content_type or ""
    file_bytes = await file.read()

    is_pdf  = "pdf" in content_type or filename.lower().endswith(".pdf")
    is_docx = "wordprocessingml" in content_type or filename.lower().endswith(".docx")
    is_doc  = filename.lower().endswith(".doc")
    is_text = "text" in content_type or filename.lower().endswith(".txt")

    if is_pdf:
        reader = PdfReader(io.BytesIO(file_bytes))
        pages  = [page.extract_text() or "" for page in reader.pages]
        text   = "\n\n".join(p for p in pages if p.strip())
    elif is_docx:
        doc  = Document(io.BytesIO(file_bytes))
        text = "\n".join(para.text for para in doc.paragraphs)
    elif is_text:
        text = file_bytes.decode("utf-8", errors="replace")
    elif is_doc:
        raise HTTPException(
            status_code=415,
            detail="Le format .doc (ancien Word) n'est pas supporté. Veuillez enregistrer le fichier en .docx et réessayer.",
        )
    else:
        raise HTTPException(
            status_code=415,
            detail=f"Format non supporté : {filename}. Formats acceptés : PDF, DOCX, TXT.",
        )

    if not text.strip():
        raise HTTPException(
            status_code=422,
            detail="Aucun texte n'a pu être extrait du document. Il est peut-être scanné ou protégé.",
        )

    return {"text": text, "filename": filename}

# ─── Generate (Ollama proxy + streaming + auto-save to history) ───────────────

@app.post("/api/generate")
async def generate(request: Request):
    body = await request.json()
    agent_id = body.get("agent")
    if not agent_id:
        raise HTTPException(status_code=400, detail="agent is required")

    try:
        prompt = build_prompt(agent_id, body)
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))

    async def stream_and_save() -> AsyncGenerator[bytes, None]:
        full_output: list[str] = []
        try:
            async with httpx.AsyncClient(timeout=120) as client:
                async with client.stream(
                    "POST",
                    f"{OLLAMA_URL}/api/generate",
                    json={"model": OLLAMA_MODEL, "prompt": prompt, "stream": True},
                ) as response:
                    async for line in response.aiter_lines():
                        if line:
                            yield (line + "\n").encode()
                            try:
                                chunk = json.loads(line)
                                if chunk.get("response"):
                                    full_output.append(chunk["response"])
                            except json.JSONDecodeError:
                                pass
        except Exception as e:
            error_line = json.dumps({"error": True, "message": str(e)})
            yield (error_line + "\n").encode()
        finally:
            output_text = "".join(full_output)
            if output_text:
                async with _history_lock:
                    history = read_json(HISTORY_FILE)
                    history.insert(0, {
                        "id": str(uuid.uuid4()),
                        "agent": agent_id,
                        "agent_label": AGENT_LABELS.get(agent_id, agent_id),
                        "input": build_input_summary(agent_id, body),
                        "fields": body.get("fields"),
                        "output": output_text,
                        "created_at": datetime.now().isoformat(),
                    })
                    write_json(HISTORY_FILE, history)

    return StreamingResponse(stream_and_save(), media_type="text/plain")

# ─── History ──────────────────────────────────────────────────────────────────

@app.get("/api/history")
def get_history(agent: Optional[str] = None, limit: int = 20):
    history = read_json(HISTORY_FILE)
    if agent:
        history = [h for h in history if h["agent"] == agent]
    return history[:limit]


@app.delete("/api/history/{entry_id}", status_code=204)
async def delete_history_entry(entry_id: str):
    async with _history_lock:
        history = read_json(HISTORY_FILE)
        updated = [h for h in history if h["id"] != entry_id]
        if len(updated) == len(history):
            raise HTTPException(status_code=404, detail="History entry not found")
        write_json(HISTORY_FILE, updated)


@app.delete("/api/history", status_code=204)
async def clear_history():
    async with _history_lock:
        write_json(HISTORY_FILE, [])

# ─── Templates ────────────────────────────────────────────────────────────────

@app.get("/api/templates")
def get_templates(agent: Optional[str] = None):
    templates = read_json(TEMPLATES_FILE)
    if agent:
        templates = [t for t in templates if t["agent"] == agent]
    return templates


@app.post("/api/templates", status_code=201)
async def create_template(request: Request):
    body = await request.json()
    agent  = body.get("agent", "").strip()
    label  = body.get("label", "").strip()
    prompt = body.get("prompt", "").strip()
    if not agent or not label or not prompt:
        raise HTTPException(status_code=400, detail="agent, label and prompt are required")
    templates = read_json(TEMPLATES_FILE)
    entry = {
        "id":         str(uuid.uuid4()),
        "agent":      agent,
        "label":      label,
        "prompt":     prompt,
        "created_at": datetime.now().isoformat(),
    }
    templates.append(entry)
    write_json(TEMPLATES_FILE, templates)
    return entry


@app.delete("/api/templates/{template_id}", status_code=204)
def delete_template(template_id: str):
    templates = read_json(TEMPLATES_FILE)
    updated = [t for t in templates if t["id"] != template_id]
    if len(updated) == len(templates):
        raise HTTPException(status_code=404, detail="Template not found")
    write_json(TEMPLATES_FILE, updated)

# ─── Feedback ─────────────────────────────────────────────────────────────────

@app.post("/api/feedback", status_code=201)
async def post_feedback(request: Request):
    body   = await request.json()
    agent  = body.get("agent", "").strip()
    rating = body.get("rating", "").strip()
    if agent not in AGENT_LABELS:
        raise HTTPException(status_code=400, detail="invalid agent")
    if rating not in ("up", "down"):
        raise HTTPException(status_code=400, detail="rating must be 'up' or 'down'")
    feedback = read_json(FEEDBACK_FILE)
    entry = {
        "id":         str(uuid.uuid4()),
        "history_id": body.get("history_id", ""),
        "agent":      agent,
        "rating":     rating,
        "created_at": datetime.now().isoformat(),
    }
    feedback.append(entry)
    write_json(FEEDBACK_FILE, feedback)
    return entry

# ─── Export ───────────────────────────────────────────────────────────────────

AGENT_LABELS_SHORT = {
    "rag":     "Consultation dossiers",
    "invoice": "Facture",
    "content": "Contenu",
}


def _make_docx(content: str, agent: str) -> bytes:
    doc = Document()
    heading = doc.add_heading("Le Play Avocats", level=1)
    heading.runs[0].font.color.rgb = RGBColor(0x0D, 0x0D, 0x0D)
    meta = doc.add_paragraph(
        f"{AGENT_LABELS_SHORT.get(agent, agent)}  —  {datetime.now().strftime('%d/%m/%Y')}"
    )
    meta.runs[0].font.size = Pt(10)
    doc.add_paragraph("─" * 60)
    for line in content.split("\n"):
        doc.add_paragraph(line)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _make_pdf(content: str, agent: str) -> bytes:
    buf = io.BytesIO()
    doc_pdf = SimpleDocTemplate(
        buf, pagesize=A4,
        leftMargin=2.5*cm, rightMargin=2.5*cm,
        topMargin=2.5*cm,  bottomMargin=2.5*cm,
    )
    styles     = getSampleStyleSheet()
    title_style = ParagraphStyle("title", parent=styles["Heading1"], fontSize=16, spaceAfter=4)
    meta_style  = ParagraphStyle("meta",  parent=styles["Normal"],  fontSize=9,
                                 textColor=(0.4, 0.4, 0.4), spaceAfter=12)
    body_style  = ParagraphStyle("body",  parent=styles["Normal"],  fontSize=11, leading=16)

    story = [
        Paragraph("Le Play Avocats", title_style),
        Paragraph(
            f"{AGENT_LABELS_SHORT.get(agent, agent)} &mdash; {datetime.now().strftime('%d/%m/%Y')}",
            meta_style,
        ),
        Spacer(1, 0.3*cm),
    ]
    for line in content.split("\n"):
        safe_line = line.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        story.append(Paragraph(safe_line or "&nbsp;", body_style))
    doc_pdf.build(story)
    return buf.getvalue()


@app.post("/api/export")
async def export_document(request: Request):
    body    = await request.json()
    content = body.get("content", "").strip()
    agent   = body.get("agent", "")
    fmt     = body.get("format", "pdf")

    if not content:
        raise HTTPException(status_code=400, detail="content is required")
    if fmt not in ("pdf", "docx"):
        raise HTTPException(status_code=400, detail="format must be 'pdf' or 'docx'")

    date_str  = datetime.now().strftime("%Y%m%d")
    label     = AGENT_LABELS_SHORT.get(agent, "document").replace(" ", "-").lower()
    filename  = f"leplay-{label}-{date_str}.{fmt}"

    if fmt == "docx":
        file_bytes   = _make_docx(content, agent)
        content_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    else:
        file_bytes   = _make_pdf(content, agent)
        content_type = "application/pdf"

    return Response(
        content=file_bytes,
        media_type=content_type,
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )

# ─── Cleaner ──────────────────────────────────────────────────────────────────

@app.post("/api/cleaner/scan")
async def cleaner_scan():
    settings = get_cleaner_settings()
    proposals, total_scanned = file_cleaner.scan_for_deletion(settings["folders"])
    total_size_bytes = sum(p["size_bytes"] for p in proposals)
    return {
        "proposals": proposals,
        "total_scanned": total_scanned,
        "total_size_bytes": total_size_bytes,
    }


@app.post("/api/cleaner/apply")
async def cleaner_apply(request: Request):
    body = await request.json()
    paths = body.get("paths", [])
    deleted, errors = [], []
    for path_str in paths:
        try:
            Path(path_str).unlink()
            deleted.append(path_str)
        except FileNotFoundError:
            errors.append({"path": path_str, "error": "Fichier introuvable"})
        except Exception as e:
            errors.append({"path": path_str, "error": str(e)})
    return {"deleted": deleted, "errors": errors}


@app.get("/api/cleaner/settings")
async def get_settings():
    return get_cleaner_settings()


@app.post("/api/cleaner/settings")
async def update_settings(request: Request):
    body = await request.json()
    settings = get_cleaner_settings()
    if "folders" in body:
        settings["folders"] = body["folders"]
    CLEANER_SETTINGS_FILE.write_text(
        json.dumps(settings, ensure_ascii=False), encoding="utf-8"
    )
    return settings


# ─── Static files (must come last so /api routes are matched first) ───────────

app.mount("/", StaticFiles(directory="interface", html=True), name="static")
